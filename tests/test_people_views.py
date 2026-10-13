"""People with their own links, time zones, and month / 3-month / year views."""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from argon2 import PasswordHasher
from conftest import COWORKER_SLUG, PARTNER_PASSWORD, PARTNER_SLUG, WEEK
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_policy import SECRETS_FROM_COWORKERS

from glimpse.config import Level
from glimpse.core import Glimpse, period_bounds
from glimpse.policy import BUSY_TITLE, VisibleEvent
from glimpse.render.periods import build_heat, busy_minutes, heat_level
from glimpse.store import ConfigStore
from glimpse.web.app import create_app
from glimpse.web.auth import feed_token
from glimpse.web.state import SeenStore, ago
from glimpse.web.viewer import person_viewer

NY = ZoneInfo("America/New_York")
DUBAI = ZoneInfo("Asia/Dubai")  # UTC+4 all year: 8 h ahead of New York in late October
W = WEEK.isoformat()
ALEX_PW = "alex-own-password"


@pytest.fixture
def store(config_path):
    return ConfigStore(config_path)


def build(store, mutate=None, today=WEEK, seen=None):
    if mutate:
        store.update(mutate)
    g = Glimpse(store.current())
    g.cache.refresh_all()
    g.today = lambda: today
    app = create_app(g.cfg, glimpse=g, background_refresh=False, seen=seen)
    return g, TestClient(app, follow_redirects=False)


def add_people(d):
    d.add_person("Alex", "partner")
    d.add_person("Sam", "coworkers")


def plink(g, pid):
    return person_viewer(g.cfg, pid).base


def login(client, base, password, nxt=W):
    return client.post(f"{base}/login", data={"password": password, "next": nxt})


# --- people --------------------------------------------------------------------------------


def test_person_link_shows_their_audience_and_moving_changes_it(store):
    g, c = build(store, add_people)
    alex = plink(g, "alex")
    assert c.get(f"{alex}/week/{W}").status_code == 303  # partner audience: password needed
    assert login(c, alex, PARTNER_PASSWORD).status_code == 303
    assert "Dinner with Sam" in c.get(f"{alex}/week/{W}").text

    # Move Alex to coworkers: same link, coworker view, immediately.
    store.update(lambda d: d.people["alex"].__setitem__("audience", "coworkers"))
    g.apply(store.current())
    html = c.get(f"{alex}/week/{W}").text
    assert "Lab work" in html
    for word in SECRETS_FROM_COWORKERS:
        assert word not in html, word


def test_removed_expired_and_switched_off_links(store):
    g, c = build(store, add_people)
    sam = plink(g, "sam")
    assert c.get(f"{sam}/week/{W}").status_code == 200

    store.update(lambda d: d.people["sam"].__setitem__("expires", (WEEK - timedelta(days=1)).isoformat()))
    g.apply(store.current())
    r = c.get(f"{sam}/week/{W}")
    assert r.status_code == 410 and "expired" in r.text
    tok = feed_token(g.cfg, person_viewer(g.cfg, "sam"))
    assert c.get(f"{sam}/feed/{tok}.ics").status_code == 404

    store.update(lambda d: d.people["sam"].update(expires=WEEK.isoformat(), enabled=False))
    g.apply(store.current())
    assert c.get(f"{sam}/week/{W}").status_code == 410

    store.update(lambda d: d.remove_person("sam"))
    g.apply(store.current())
    assert c.get(f"{sam}/week/{W}").status_code == 404


def test_expiry_day_itself_still_works(store):
    g, c = build(store, lambda d: d.add_person("Sam", "coworkers", WEEK.isoformat()))
    assert c.get(f"{plink(g, 'sam')}/week/{W}").status_code == 200


def test_person_password_modes(store):
    h = PasswordHasher().hash(ALEX_PW)

    def setup(d):
        d.add_person("Alex", "partner")
        d.set_person_password("alex", h)
        d.add_person("Jo", "partner")
        d.people["jo"]["use_audience_password"] = False

    g, c = build(store, setup)
    alex, jo = plink(g, "alex"), plink(g, "jo")
    assert login(c, alex, PARTNER_PASSWORD).status_code == 401  # their own password replaces it
    assert login(c, alex, ALEX_PW).status_code == 303
    assert c.get(f"{alex}/week/{W}").status_code == 200
    assert c.get(f"{jo}/week/{W}").status_code == 200  # no password at all


def test_cookies_dont_cross_between_links(store):
    g, c = build(store, add_people)
    login(c, f"/a/{PARTNER_SLUG}", PARTNER_PASSWORD)
    assert c.get(f"/a/{PARTNER_SLUG}/week/{W}").status_code == 200
    assert c.get(f"{plink(g, 'alex')}/week/{W}").status_code == 303


def test_rotating_a_person_link_only_affects_them(store):
    g, c = build(store, add_people)
    old = plink(g, "alex")
    store.update(lambda d: d.rotate_person_slug("alex"))
    g.apply(store.current())
    assert c.get(f"{old}/week/{W}").status_code == 404
    assert c.get(f"{plink(g, 'sam')}/week/{W}").status_code == 200
    assert c.get(f"/a/{COWORKER_SLUG}/week/{W}").status_code == 200


def test_shared_link_can_be_switched_off(store):
    def setup(d):
        add_people(d)
        d.audiences["coworkers"]["shared_link"] = False

    g, c = build(store, setup)
    assert c.get(f"/a/{COWORKER_SLUG}/week/{W}").status_code == 404
    assert c.get(f"/a/{COWORKER_SLUG}/feed/{feed_token(g.cfg, 'coworkers')}.ics").status_code == 404
    assert c.get(f"{plink(g, 'sam')}/week/{W}").status_code == 200


def test_person_feed_works_and_is_personal(store):
    g, c = build(store, add_people)
    v = person_viewer(g.cfg, "sam")
    tok = feed_token(g.cfg, v)
    assert tok != feed_token(g.cfg, "coworkers")
    r = c.get(f"{v.base}/feed/{tok}.ics")
    assert r.status_code == 200 and b"Lab work" in r.content
    assert c.get(f"/a/{COWORKER_SLUG}/feed/{tok}.ics").status_code == 404


def test_deleting_an_audience_removes_its_people(store):
    store.update(add_people)
    cfg = store.update(lambda d: d.remove_audience("coworkers"))
    assert list(cfg.people) == ["alex"]


def test_people_must_reference_real_audiences(store):
    from glimpse.config import ConfigError

    with pytest.raises(ConfigError, match="not defined"):
        store.update(lambda d: d.add_person("Ghost", "nope"))


def test_last_viewed_is_recorded_in_state_json(store, tmp_path):
    seen = SeenStore(tmp_path / "state.json")
    g, c = build(store, add_people, seen=seen)
    assert seen.get("person:sam") is None
    c.get(f"{plink(g, 'sam')}/week/{W}")
    c.get(f"/a/{COWORKER_SLUG}/week/{W}")
    assert seen.get("person:sam") is not None and seen.get("audience:coworkers") is not None
    seen.flush()
    data = json.loads((tmp_path / "state.json").read_text())
    assert set(data["last_seen"]) >= {"person:sam", "audience:coworkers"}
    assert SeenStore(tmp_path / "state.json").get("person:sam") == seen.get("person:sam")
    # nothing about people is written to config.yaml by viewing
    assert "last_seen" not in store.config_path.read_text()


def test_ago():
    now = datetime(2026, 10, 26, 12, tzinfo=NY)
    assert ago(None) == "never"
    assert ago(now - timedelta(seconds=30), now) == "just now"
    assert ago(now - timedelta(minutes=5), now) == "5 min ago"
    assert ago(now - timedelta(hours=3), now) == "3 h ago"
    assert ago(now - timedelta(days=1, hours=2), now) == "yesterday"
    assert ago(now - timedelta(days=4), now) == "4 days ago"


# --- time zones ----------------------------------------------------------------------------


def find(evs, title):
    return next(e for e in evs if e.title == title)


def test_viewer_ahead_sees_late_events_on_the_next_day(glimpse):
    partner = glimpse.cfg.audiences["partner"]
    evs = glimpse.week(partner, WEEK, DUBAI)
    movie = find(evs, "Late movie")  # Fri 22:30 New York
    assert movie.start.tzinfo == DUBAI
    assert (movie.start.weekday(), movie.start.hour, movie.start.minute) == (5, 6, 30)  # Sat 06:30


def test_all_day_events_keep_their_date(glimpse):
    partner = glimpse.cfg.audiences["partner"]
    bday = find(glimpse.week(partner, WEEK, DUBAI), "Mom's birthday")
    assert bday.all_day and bday.start.date() == date(2026, 10, 30)


def test_hour_limits_stay_on_the_owners_clock(glimpse):
    cw = glimpse.cfg.audiences["coworkers"]
    evs = glimpse.week(cw, WEEK, DUBAI)
    seminar = find(evs, "Evening seminar")  # 17:30-19:00 NY, clipped at 18:00 NY
    assert seminar.end - seminar.start == timedelta(minutes=30)
    assert seminar.start.astimezone(NY).hour == 17
    lab = find(evs, "Lab work")
    assert lab.start.hour == 17  # 09:00 NY = 17:00 Dubai


def test_dst_change_in_owner_zone(glimpse):
    partner = glimpse.cfg.audiences["partner"]
    utc = ZoneInfo("UTC")
    before = [e for e in glimpse.week(partner, WEEK, utc) if e.title == "Lab work"]
    after = [e for e in glimpse.week(partner, WEEK + timedelta(weeks=1), utc) if e.title == "Lab work"]
    assert before[0].start.hour == 13 and after[0].start.hour == 14  # 9:00 EDT, then 9:00 EST


def test_zone_switch_sets_cookie_and_page_follows(glimpse):
    c = TestClient(create_app(glimpse.cfg, glimpse=glimpse, background_refresh=False), follow_redirects=False)
    r = c.get(f"/a/{COWORKER_SLUG}/week/{W}?tz=Asia/Dubai")
    assert r.status_code == 303 and r.headers["location"] == f"/a/{COWORKER_SLUG}/week/{W}"
    assert re.match(r'glimpse_tz="?Asia/Dubai"?;', r.headers["set-cookie"])
    html = c.get(f"/a/{COWORKER_SLUG}/week/{W}").text
    assert "Times in <b>Dubai" in html and "?tz=owner" in html
    c.get(f"/a/{COWORKER_SLUG}/week/{W}?tz=owner")
    assert "Times in <b>New York" in c.get(f"/a/{COWORKER_SLUG}/week/{W}").text
    r = c.get(f"/a/{COWORKER_SLUG}/week/{W}?tz=Not/AZone")
    assert "set-cookie" not in r.headers


def test_audience_default_zone(store):
    g, c = build(store, lambda d: d.audiences["coworkers"].__setitem__("timezone", "Asia/Dubai"))
    assert "Times in <b>Dubai" in c.get(f"/a/{COWORKER_SLUG}/week/{W}").text


def test_owner_name_labels_owner_time(store):
    g, c = build(store, lambda d: d.raw.__setitem__("owner_name", "Brandon"))
    c.get(f"/a/{COWORKER_SLUG}/week/{W}?tz=Asia/Dubai")
    assert "Show Brandon&#39;s time" in c.get(f"/a/{COWORKER_SLUG}/week/{W}").text


# --- month, 3-month and year views ---------------------------------------------------------


def all_views(d):
    for aud in d.audiences.values():
        aud["views"] = ["week", "month", "quarter", "year"]


def test_views_not_enabled_are_404(store):
    g, c = build(store, lambda d: d.audiences["coworkers"].__setitem__("views", ["week"]))
    assert c.get(f"/a/{COWORKER_SLUG}/month/2026-10-01").status_code == 404
    assert 'class="views"' not in c.get(f"/a/{COWORKER_SLUG}/week/{W}").text  # one view: no switcher


def test_new_audiences_get_every_view(store):
    cfg = store.update(lambda d: d.add_audience("Lab group"))
    assert cfg.audiences["lab_group"].views == ["week", "month", "quarter", "year"]


def test_month_view(store):
    g, c = build(store, all_views)
    login(c, f"/a/{PARTNER_SLUG}", PARTNER_PASSWORD)
    r = c.get(f"/a/{PARTNER_SLUG}/month/2026-10-15")
    assert r.status_code == 303 and r.headers["location"].endswith("/month/2026-10-01")
    html = c.get(f"/a/{PARTNER_SLUG}/month/2026-10-01").text
    assert "October 2026" in html and "Saturday experiment" in html
    assert re.search(r"\+\d+ more", html)  # busy days are summarised
    assert f'href="/a/{PARTNER_SLUG}/week/2026-10-28"' in html  # days open their week


def test_month_view_for_coworkers_leaks_nothing(store):
    g, c = build(store, all_views)
    html = c.get(f"/a/{COWORKER_SLUG}/month/2026-10-01").text
    assert "Lab work" in html and "Busy" in html
    for word in SECRETS_FROM_COWORKERS:
        assert word not in html, word


@pytest.mark.parametrize("kind", ["quarter", "year"])
def test_heat_maps_contain_no_event_text(store, kind):
    g, c = build(store, all_views)
    first, _ = period_bounds(kind, WEEK)
    login(c, f"/a/{PARTNER_SLUG}", PARTNER_PASSWORD)
    for slug in (PARTNER_SLUG, COWORKER_SLUG):
        html = c.get(f"/a/{slug}/{kind}/{first}").text
        assert 'class="hd h' in html
        for word in ["Dinner with Sam", "Lab work", "Polymer", "Luigi", "Doctor", "Grandma", "Busy"]:
            assert word not in html, (slug, word)


def test_browse_limits_apply_to_longer_views(store):
    def setup(d):
        all_views(d)
        d.audiences["coworkers"].update(weeks_ahead=2, weeks_behind=0)

    g, c = build(store, setup)
    r = c.get(f"/a/{COWORKER_SLUG}/month/2027-03-01")  # far beyond: nearest allowed month
    assert r.status_code == 303 and r.headers["location"].endswith("/month/2026-11-01")
    html = c.get(f"/a/{COWORKER_SLUG}/month/2026-11-01").text
    # Allowed up to Sun Nov 15 (2 weeks ahead of the week of Oct 26). Later days are locked, empty.
    assert html.count("mcell locked") + html.count("mcell locked today") >= 15
    late = html.split('aria-label="Week of Mon Nov 16"')
    assert len(late) == 1  # no link into weeks they can't see
    year = c.get(f"/a/{COWORKER_SLUG}/year/2026-01-01").text
    assert 'title="Fri Dec 25: not shared"' in year
    assert 'title="Mon Oct 19: not shared"' in year  # before this week (weeks_behind 0)


def test_month_excel(store):
    g, c = build(store, all_views)
    r = c.get(f"/a/{COWORKER_SLUG}/excel-month/2026-10-01")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Month", "Events", "Busy hours"]
    text = " ".join(str(x.value) for ws in wb for row in ws.iter_rows() for x in row if x.value)
    assert "Lab work" in text
    for word in SECRETS_FROM_COWORKERS:
        assert word not in text, word


def test_default_view_and_switcher(store):
    def setup(d):
        d.audiences["coworkers"].update(views=["month", "year"], default_view="month")

    g, c = build(store, setup)
    r = c.get(f"/a/{COWORKER_SLUG}")
    assert r.status_code == 303 and "/month/" in r.headers["location"]
    html = c.get(f"/a/{COWORKER_SLUG}/month/2026-10-01").text
    assert ">Year</a>" in html and ">Week<" not in html
    assert c.get(f"/a/{COWORKER_SLUG}/week/{W}").status_code == 404


# --- building blocks -----------------------------------------------------------------------


def ve(start, end, title="x", level=Level.TITLE, all_day=False, cal="personal"):
    return VisibleEvent(
        key=title + str(start),
        start=start,
        end=end,
        all_day=all_day,
        level=level,
        title=title,
        calendar_id=None if level is Level.BUSY else cal,
        calendar_name=None if level is Level.BUSY else "P",
    )


def test_busy_minutes_counts_overlaps_once():
    d = date(2026, 10, 26)
    t = lambda h, m=0: datetime(2026, 10, 26, h, m, tzinfo=NY)  # noqa: E731
    evs = [ve(t(9), t(11)), ve(t(10), t(12)), ve(t(14), t(14, 30)), ve(t(23), t(23) + timedelta(hours=3))]
    assert busy_minutes(evs, d, NY) == 3 * 60 + 30 + 60
    assert [heat_level(m) for m in (0, 30, 120, 300, 600)] == [0, 1, 2, 3, 4]


def test_heat_spans_stitch_multi_day_items(cfg):
    days = [datetime(2026, 11, d, tzinfo=NY) for d in (10, 11, 12)]
    trip = [ve(s, s + timedelta(days=1), "Conference", all_day=True) for s in days]
    busy = [ve(s, s + timedelta(days=1), BUSY_TITLE, Level.BUSY, all_day=True) for s in days[:2]]
    one = [ve(days[0], days[0] + timedelta(days=1), "Birthday", all_day=True)]
    view = build_heat(
        cfg,
        cfg.audiences["partner"],
        date(2026, 11, 1),
        date(2026, 12, 1),
        "quarter",
        trip + busy + one,
        WEEK,
        (date(2026, 1, 1), date(2027, 1, 1)),
        NY,
    )
    spans = [(s.title, s.first.day, s.last.day) for s in view.spans]
    assert ("Conference", 10, 12) in spans and (BUSY_TITLE, 10, 11) in spans
    assert all(t != "Birthday" for t, *_ in spans)  # single days aren't spans


def test_month_chip_for_event_past_midnight(store):
    g, c = build(store, all_views)
    login(c, f"/a/{PARTNER_SLUG}", PARTNER_PASSWORD)
    html = c.get(f"/a/{PARTNER_SLUG}/month/2026-10-01").text
    sat = html.split('aria-label="Week of Sat Oct 31"')[1].split("mcell")[0]
    assert "until 1a" in sat and "Late movie" in sat
