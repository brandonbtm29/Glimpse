"""Calendar groups, keyword privacy rules and duplicate merging."""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from conftest import PARTNER_PASSWORD, PARTNER_SLUG, WEEK
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from kalends.config import ConfigError, Level, RuleCfg, build_config
from kalends.core import Kalends
from kalends.events import Event
from kalends.policy import BUSY_TITLE, cluster_duplicates, fold, rule_matches, visible_events
from kalends.store import ConfigStore
from kalends.web.app import create_app
from kalends.web.auth import feed_token

TZ = ZoneInfo("America/New_York")
SECRET = "s" * 40


def make_cfg(**extra):
    raw = {
        "config_version": 2,
        "timezone": "America/New_York",
        "server": {"secret_key": SECRET},
        "calendars": {
            "personal": {"name": "Personal", "url": "https://x.example/p.ics"},
            "work": {"name": "Work", "url": "https://x.example/w.ics"},
            "uni": {"name": "University", "url": "https://x.example/u.ics", "label": "Uni"},
            "family": {"name": "Family", "url": "https://x.example/f.ics"},
        },
        "audiences": {
            "partner": {"title": "Partner", "slug": "p" * 20, "default": "full"},
            "team": {"title": "Team", "slug": "t" * 20, "default": "busy"},
        },
    }
    for key, value in extra.items():
        if isinstance(value, dict) and isinstance(raw.get(key), dict):
            for k, v in value.items():
                if isinstance(v, dict) and isinstance(raw[key].get(k), dict):
                    raw[key][k] = {**raw[key][k], **v}
                else:
                    raw[key][k] = v
        else:
            raw[key] = value
    return build_config(raw, {}, Path("."))


def ev(cal, title, hour=10, minutes=60, uid=None, notes="", location="", private=False, free=False, day=0):
    start = datetime(2026, 10, 26 + day, hour, tzinfo=TZ)
    return Event(
        calendar_id=cal,
        uid=uid or f"{cal}-{title}-{hour}@x",
        start=start,
        end=start + timedelta(minutes=minutes),
        all_day=False,
        summary=title,
        location=location,
        description=notes,
        private=private,
        free=free,
    )


RANGE = (datetime(2026, 10, 26, tzinfo=TZ), datetime(2026, 11, 2, tzinfo=TZ))


def show(cfg, audience, events):
    return visible_events(cfg, cfg.audiences[audience], events, *RANGE)


# --- groups -------------------------------------------------------------------------------


def grouped(**aud_overrides):
    return make_cfg(
        groups={"job": {"name": "Job", "calendars": ["work", "uni"]}},
        audiences={"team": {"groups": {"job": "title"}, **aud_overrides}},
    )


def test_group_level_applies_to_members():
    cfg = grouped()
    team = cfg.audiences["team"]
    assert cfg.level_for(team, "work") is Level.TITLE
    assert cfg.level_for(team, "uni") is Level.TITLE
    assert cfg.level_for(team, "personal") is Level.BUSY  # not in the group: default
    assert cfg.level_source(team, "work") == "group"
    assert cfg.level_source(team, "personal") == "default"


def test_calendar_setting_beats_group():
    cfg = grouped(calendars={"uni": "label"})
    team = cfg.audiences["team"]
    assert cfg.level_for(team, "uni") is Level.LABEL
    assert cfg.level_source(team, "uni") == "calendar"
    assert cfg.level_for(team, "work") is Level.TITLE


def test_audience_without_group_setting_uses_default():
    cfg = grouped()
    assert cfg.level_for(cfg.audiences["partner"], "work") is Level.FULL
    assert cfg.level_source(cfg.audiences["partner"], "work") == "default"


def test_group_level_reaches_the_rendered_view():
    cfg = grouped()
    out = show(cfg, "team", [ev("work", "Standup"), ev("personal", "Dentist", hour=14)])
    assert [e.title for e in out] == ["Standup", BUSY_TITLE]


@pytest.mark.parametrize(
    ("extra", "msg"),
    [
        ({"groups": {"a": {"name": "A", "calendars": ["nope"]}}}, "nope"),
        (
            {
                "groups": {
                    "a": {"name": "A", "calendars": ["work"]},
                    "b": {"name": "B", "calendars": ["work"]},
                }
            },
            "two groups",
        ),
        ({"audiences": {"team": {"groups": {"ghost": "full"}}}}, "ghost"),
        ({"rules": [{"match": ["x"], "calendars": ["nope"]}]}, "nope"),
        ({"rules": [{"match": [" "]}]}, "keyword"),
        ({"rules": [{"match": ["x"], "fields": []}]}, "at least one"),
    ],
)
def test_invalid_groups_and_rules_rejected(extra, msg):
    with pytest.raises((ConfigError, ValueError)) as exc:
        make_cfg(**extra)
    assert msg in str(exc.value)


# --- keyword rules -------------------------------------------------------------------------


def rule(**kw):
    return RuleCfg.model_validate({"match": ["doctor"], **kw})


@pytest.mark.parametrize(
    ("words", "text", "hit"),
    [
        (["doctor"], "Doctor appointment", True),
        (["doctor"], "DOCTOR", True),
        (["dr"], "Dr. Lee", True),
        (["dr"], "Drive to airport", False),
        (["therapy"], "Thérapie", False),
        (["thérapie"], "THERAPIE session", True),
        (["cafe"], "Café with Ana", True),
        (["job interview"], "Job   interview at Acme", True),
        (["job interview"], "Job, then interview", False),
        (["c++"], "Learn C++ tonight", True),
    ],
)
def test_keyword_matching(words, text, hit):
    assert rule_matches(rule(match=words), ev("personal", text)) is hit


def test_fold():
    assert fold("Ñandú STRASSE") == "nandu strasse"


def test_rule_fields_and_calendars():
    e = ev("personal", "Appointment", notes="see the doctor", location="Clinic")
    assert rule_matches(rule(), e)
    assert not rule_matches(rule(fields=["title", "location"]), e)
    assert rule_matches(rule(match=["clinic"], fields=["location"]), e)
    assert not rule_matches(rule(calendars=["work"]), e)
    assert rule_matches(rule(calendars=["personal", "work"]), e)
    assert not rule_matches(rule(enabled=False), e)


def test_rule_turns_full_into_busy_for_everyone():
    cfg = make_cfg(rules=[{"match": ["doctor"]}])
    events = [ev("personal", "Doctor visit", location="Clinic Main St"), ev("personal", "Lunch", hour=12)]
    partner = show(cfg, "partner", events)
    assert [e.title for e in partner] == [BUSY_TITLE, "Lunch"]
    assert all(e.location in ("", None) for e in partner if e.title == BUSY_TITLE)


def test_hidden_rule_removes_event():
    cfg = make_cfg(rules=[{"match": ["interview"], "action": "hidden"}])
    out = show(cfg, "partner", [ev("work", "Interview at Acme")])
    assert out == []


def test_rule_never_increases_detail():
    cfg = make_cfg(
        rules=[{"match": ["doctor"], "action": "busy"}],
        audiences={"team": {"calendars": {"family": "hidden"}}},
    )
    assert show(cfg, "team", [ev("family", "Doctor with Mom")]) == []


def test_rules_combine_most_restrictive_wins():
    cfg = make_cfg(rules=[{"match": ["doctor"]}, {"match": ["mom"], "action": "hidden"}])
    assert show(cfg, "partner", [ev("family", "Doctor with Mom")]) == []


# --- duplicate merging ----------------------------------------------------------------------


def test_same_uid_on_two_calendars_shows_once_at_best_level():
    cfg = make_cfg(audiences={"team": {"calendars": {"work": "title"}}})
    events = [ev("personal", "Project sync", uid="abc@x"), ev("work", "Project sync", uid="abc@x")]
    out = show(cfg, "team", events)
    assert [(e.title, e.calendar_id, e.copies) for e in out] == [("Project sync", "work", 2)]


def test_title_and_time_match_across_calendars():
    cfg = make_cfg()
    events = [ev("personal", "Dentist "), ev("family", "dentist", uid="other@x")]
    assert len(show(cfg, "partner", events)) == 1


def test_near_misses_are_not_merged():
    cfg = make_cfg()
    events = [
        ev("personal", "Dentist"),
        ev("family", "Dentist", minutes=90),  # different end
        ev("work", "Dentist", hour=11),  # different start
        ev("uni", "Dentists"),  # different title
    ]
    assert len(show(cfg, "partner", events)) == 4


def test_same_title_on_same_calendar_is_two_events():
    cfg = make_cfg()
    events = [ev("work", "Office hours", uid="a@x"), ev("work", "Office hours", uid="b@x")]
    assert len(show(cfg, "partner", events)) == 2


def test_restriction_on_any_copy_wins():
    cfg = make_cfg()
    events = [
        ev("work", "Review", uid="r@x"),
        ev("personal", "Review", uid="r@x", notes="#private"),
    ]
    out = show(cfg, "partner", events)
    assert [e.title for e in out] == [BUSY_TITLE]


def test_rule_on_any_copy_wins():
    cfg = make_cfg(rules=[{"match": ["salary"], "fields": ["notes"]}])
    events = [ev("work", "1:1", uid="o@x"), ev("personal", "1:1", uid="o@x", notes="ask about salary")]
    assert [e.title for e in show(cfg, "partner", events)] == [BUSY_TITLE]


def test_hidden_copy_does_not_hide_a_visible_copy():
    cfg = make_cfg(audiences={"team": {"calendars": {"family": "hidden", "work": "title"}}})
    events = [ev("family", "Flight", uid="f@x"), ev("work", "Flight", uid="f@x")]
    assert [e.title for e in show(cfg, "team", events)] == ["Flight"]


def test_free_copy_does_not_make_busy_copy_free():
    cfg = make_cfg()
    clusters = cluster_duplicates(
        [ev("work", "Trip", uid="t@x", free=True), ev("personal", "Trip", uid="t@x")]
    )
    assert len(clusters) == 1
    best, _ = clusters[0].resolve(cfg, cfg.audiences["team"])
    assert best.free is False


def test_merging_can_be_turned_off():
    cfg = make_cfg(overrides={"merge_duplicates": False})
    events = [ev("personal", "Sync", uid="s@x"), ev("work", "Sync", uid="s@x")]
    out = show(cfg, "partner", events)
    assert len(out) == 2 and all(e.copies == 1 for e in out)


def test_recurring_instances_with_shared_uid_stay_separate():
    cfg = make_cfg()
    events = [ev("work", "Standup", uid="daily@x", day=d) for d in range(5)]
    events += [ev("personal", "Standup", uid="daily@x", day=d) for d in range(5)]
    out = show(cfg, "partner", events)
    assert len(out) == 5 and all(e.copies == 2 for e in out)


def test_busy_duplicates_merge_into_one_block():
    cfg = make_cfg()
    events = [ev("personal", "Secret", uid="z@x"), ev("work", "Secret", uid="z@x")]
    out = show(cfg, "team", events)
    assert [(e.title, e.start.hour, e.end.hour) for e in out] == [(BUSY_TITLE, 10, 11)]


# --- no leaks through any output, using the real fixtures ------------------------------------


@pytest.fixture
def ruled(config_path):
    store = ConfigStore(config_path)
    store.update(lambda d: d.set_rule(None, {"match": ["luigi's"], "fields": ["location"], "action": "busy"}))
    g = Kalends(store.current())
    g.cache.refresh_all()
    return g


def test_rule_holds_across_html_excel_and_ics(ruled):
    client = TestClient(
        create_app(ruled.cfg, kalends=ruled, background_refresh=False), follow_redirects=False
    )
    w = WEEK.isoformat()
    client.post(f"/a/{PARTNER_SLUG}/login", data={"password": PARTNER_PASSWORD, "next": w})
    html = client.get(f"/a/{PARTNER_SLUG}/week/{w}").text
    assert "Lab work" in html  # partner still sees normal events
    excel = client.get(f"/a/{PARTNER_SLUG}/excel/{w}").content
    wb = load_workbook(io.BytesIO(excel))
    xl = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    ics = client.get(f"/a/{PARTNER_SLUG}/feed/{feed_token(ruled.cfg, 'partner')}.ics").text
    for out in (html, xl, ics):
        assert "Dinner with Sam" not in out
        assert "Luigi" not in out


# --- store ------------------------------------------------------------------------------------


def test_store_groups_and_cleanup(config_path):
    store = ConfigStore(config_path)

    def setup(d):
        gid = d.add_group("Work stuff")
        d.set_calendar_group("lab", gid)
        d.set_calendar_group("classes", gid)
        d.audiences["coworkers"]["groups"] = {gid: "title"}
        d.set_rule(None, {"match": ["exam"], "calendars": ["classes"]})

    cfg = store.update(setup)
    assert cfg.groups["work_stuff"].calendars == ["lab", "classes"]
    assert cfg.group_of("classes") == "work_stuff"

    cfg = store.update(lambda d: d.set_calendar_group("lab", None))
    assert cfg.groups["work_stuff"].calendars == ["classes"]

    cfg = store.update(lambda d: d.remove_calendar("classes"))
    assert cfg.groups["work_stuff"].calendars == []
    assert cfg.rules[0].enabled is False and cfg.rules[0].calendars == []  # not widened to all

    cfg = store.update(lambda d: d.remove_group("work_stuff"))
    assert cfg.groups == {} and cfg.audiences["coworkers"].groups == {}


def test_calendar_moves_between_groups(config_path):
    store = ConfigStore(config_path)

    def setup(d):
        a, b = d.add_group("A"), d.add_group("B")
        d.set_calendar_group("lab", a)
        d.set_calendar_group("lab", b)

    cfg = store.update(setup)
    assert cfg.groups["a"].calendars == [] and cfg.groups["b"].calendars == ["lab"]
