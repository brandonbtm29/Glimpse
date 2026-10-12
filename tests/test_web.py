import io

import icalendar
import pytest
from conftest import COWORKER_SLUG, PARTNER_PASSWORD, PARTNER_SLUG, WEEK
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_policy import SECRETS_FROM_COWORKERS

from glimpse.web.app import create_app
from glimpse.web.auth import feed_token

W = WEEK.isoformat()


@pytest.fixture
def client(glimpse):
    app = create_app(glimpse.cfg, glimpse=glimpse, background_refresh=False)
    return TestClient(app, follow_redirects=False)


def login(client, password=PARTNER_PASSWORD):
    return client.post(f"/a/{PARTNER_SLUG}/login", data={"password": password, "next": W})


# --- routing and 404s ------------------------------------------------------------


def test_root_and_unknown_slug_are_404(client):
    assert client.get("/").status_code == 404
    assert client.get("/a/notARealSlug_0000000000/week/" + W).status_code == 404


def test_home_redirects_to_this_week(client):
    r = client.get(f"/a/{COWORKER_SLUG}")
    assert r.status_code == 303 and "/week/" in r.headers["location"]


def test_mid_week_date_redirects_to_week_start(client):
    r = client.get(f"/a/{COWORKER_SLUG}/week/2026-10-29")
    assert r.status_code == 303 and r.headers["location"].endswith("/week/2026-10-26")


def test_out_of_range_week_is_clamped(client, glimpse):
    glimpse.cfg.audiences["coworkers"].weeks_ahead = 1
    r = client.get(f"/a/{COWORKER_SLUG}/week/2030-01-07")
    assert r.status_code == 303


def test_bad_date_is_404(client):
    assert client.get(f"/a/{COWORKER_SLUG}/week/not-a-date").status_code == 404
    assert client.get(f"/a/{COWORKER_SLUG}/week/2026-13-40").status_code == 404


# --- link-only audience ---------------------------------------------------------------


def test_coworker_page_renders_without_password(client):
    r = client.get(f"/a/{COWORKER_SLUG}/week/{W}")
    assert r.status_code == 200
    assert "Lab work" in r.text and "Class" in r.text and "Busy" in r.text


def test_coworker_page_never_contains_hidden_text(client):
    html = client.get(f"/a/{COWORKER_SLUG}/week/{W}").text
    for word in SECRETS_FROM_COWORKERS:
        assert word not in html, word
    # weekends aren't columns for a Mon-Fri audience
    assert ">Sat<" not in html and ">Sun<" not in html


def test_security_headers(client):
    r = client.get(f"/a/{COWORKER_SLUG}/week/{W}")
    csp = r.headers["content-security-policy"]
    assert "default-src 'none'" in csp and "script-src 'self';" in csp  # own small script only
    assert "unsafe-inline" not in csp.split("script-src")[1].split(";")[0]
    assert r.headers["referrer-policy"] == "no-referrer"
    assert "noindex" in r.headers["x-robots-tag"]
    assert r.headers["cache-control"] == "no-store"
    assert r.headers["x-frame-options"] == "DENY"


# --- password audience -------------------------------------------------------------------


def test_partner_requires_login(client):
    r = client.get(f"/a/{PARTNER_SLUG}/week/{W}")
    assert r.status_code == 303 and "/login" in r.headers["location"]
    page = client.get(r.headers["location"])
    assert page.status_code == 200 and 'type="password"' in page.text
    assert "Dinner" not in page.text


def test_wrong_password_rejected(client):
    r = login(client, "nope")
    assert r.status_code == 401 and 'role="alert"' in r.text
    assert not r.cookies


def test_login_sets_cookie_and_shows_details(client):
    r = login(client)
    assert r.status_code == 303 and r.headers["location"].endswith(f"/week/{W}")
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=lax" in cookie and f"path=/a/{PARTNER_SLUG}".lower() in cookie
    page = client.get(f"/a/{PARTNER_SLUG}/week/{W}")
    assert page.status_code == 200
    assert "Dinner with Sam" in page.text and "Luigi" in page.text
    assert "Surprise" not in page.text and "Doctor" not in page.text


def test_cookie_for_one_audience_does_not_open_another(client, glimpse):
    from glimpse.web.auth import cookie_name, make_session

    glimpse.cfg.audiences["coworkers"].password_hash = glimpse.cfg.audiences["partner"].password_hash
    client.cookies.set(cookie_name("coworkers"), make_session(glimpse.cfg, "partner"))
    r = client.get(f"/a/{COWORKER_SLUG}/week/{W}")
    assert r.status_code == 303


def test_tampered_cookie_rejected(client):
    from glimpse.web.auth import cookie_name

    client.cookies.set(cookie_name("partner"), "eyJhIjoicGFydG5lciJ9.forged.sig")
    assert client.get(f"/a/{PARTNER_SLUG}/week/{W}").status_code == 303


def test_changing_password_logs_everyone_out(client, glimpse):
    login(client)
    assert client.get(f"/a/{PARTNER_SLUG}/week/{W}").status_code == 200
    from argon2 import PasswordHasher

    glimpse.cfg.audiences["partner"].password_hash = PasswordHasher().hash("a new password")
    assert client.get(f"/a/{PARTNER_SLUG}/week/{W}").status_code == 303


def test_rate_limit_blocks_after_five_failures(client):
    for _ in range(5):
        assert login(client, "wrong").status_code == 401
    r = login(client)  # even the right password is refused while blocked
    assert r.status_code == 429


def test_next_param_cannot_redirect_offsite(client):
    r = client.post(
        f"/a/{PARTNER_SLUG}/login", data={"password": PARTNER_PASSWORD, "next": "https://evil.example"}
    )
    assert r.status_code == 303 and r.headers["location"] == f"/a/{PARTNER_SLUG}"


def test_logout_clears_cookie(client):
    login(client)
    r = client.get(f"/a/{PARTNER_SLUG}/logout")
    assert r.status_code == 303
    client.cookies.clear()
    assert client.get(f"/a/{PARTNER_SLUG}/week/{W}").status_code == 303


# --- exports ----------------------------------------------------------------------------------


def test_excel_export_for_coworkers(client):
    r = client.get(f"/a/{COWORKER_SLUG}/excel/{W}")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Week", "Events"]
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Lab work" in text and "Busy" in text
    for word in SECRETS_FROM_COWORKERS:
        assert word not in text, word


def test_excel_requires_login_for_partner(client):
    assert client.get(f"/a/{PARTNER_SLUG}/excel/{W}").status_code == 303
    login(client)
    assert client.get(f"/a/{PARTNER_SLUG}/excel/{W}").status_code == 200


def test_ics_feed_needs_token(client, glimpse):
    assert client.get(f"/a/{COWORKER_SLUG}/feed/wrongtoken.ics").status_code == 404
    tok = feed_token(glimpse.cfg, "coworkers")
    r = client.get(f"/a/{COWORKER_SLUG}/feed/{tok}.ics")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/calendar")
    cal = icalendar.Calendar.from_ical(r.content)
    summaries = {str(e.get("summary")) for e in cal.walk("VEVENT")}
    assert {"Lab work", "Busy", "Class"} <= summaries
    for word in SECRETS_FROM_COWORKERS:
        assert word.encode() not in r.content, word
    assert b"@test" not in r.content  # original UIDs are not exposed


def test_feed_token_is_audience_specific(glimpse):
    assert feed_token(glimpse.cfg, "partner") != feed_token(glimpse.cfg, "coworkers")


def test_healthz_reveals_no_names(client):
    r = client.get("/healthz").json()
    assert r == {"ok": True, "feeds_ok": 4, "feeds_total": 4}


def test_event_text_is_html_escaped(client, glimpse):
    evil = (
        "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nUID:x@test\r\n"
        "DTSTART;TZID=America/New_York:20261026T130000\r\n"
        "DTEND;TZID=America/New_York:20261026T140000\r\n"
        'SUMMARY:<script>alert(1)</script>"><img src=x onerror=alert(2)>\r\n'
        "LOCATION:<b>loc</b>\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"
    )
    st = glimpse.cache.states["personal"]
    st.text, st.version = evil, st.version + 1
    login(client)
    html = client.get(f"/a/{PARTNER_SLUG}/week/{W}").text
    assert "<script>alert" not in html and "<img src=x" not in html and "<b>loc</b>" not in html
    assert "&lt;script&gt;" in html
